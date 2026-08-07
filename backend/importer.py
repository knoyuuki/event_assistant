"""Import personnel data from the Excel directory and match photos by name."""

import os
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Optional

import openpyxl

from database import get_connection, normalize_dept_name, DEPT_NAME_ALIASES
from models import POSITION_LEVEL_MAP, PERSON_TYPES

PEOPLE_DIR = str(Path(__file__).resolve().parent.parent / "people")
PHOTOS_DIR = str(Path(__file__).resolve().parent.parent / "photos")
EXCEL_FILE = "企业通讯录人员名单.xlsx"
MARKDOWN_FILE = "企业通讯录部门层级.md"

# Supported image extensions
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}

# Threshold for fuzzy department matching
FUZZY_THRESHOLD = 0.5


def _load_db_departments() -> list[str]:
    """Load all department names from the database."""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT name FROM departments ORDER BY sort_order, id")
            return [row["name"] for row in cur.fetchall()]
    finally:
        conn.close()


def _fuzzy_match_dept(name: str, dept_list: list[str], threshold: float = FUZZY_THRESHOLD) -> Optional[str]:
    """
    Find the best matching department name from dept_list.
    Uses substring match (high priority) then SequenceMatcher ratio.
    Returns None if no match above threshold.
    """
    if not name or not dept_list:
        return None

    best_score = 0.0
    best_match = None

    for dept in dept_list:
        # Substring/contains match gets high confidence
        if name in dept or dept in name:
            score = 0.85
        else:
            score = SequenceMatcher(None, name, dept).ratio()

        if score > best_score and score >= threshold:
            best_score = score
            best_match = dept

    return best_match


def resolve_department_fuzzy(
    dept1: str,
    dept2: str,
    db_depts: list[str],
    aliases: dict[str, str],
) -> tuple[str, str]:
    """
    Resolve a person's department from Excel dept1/dept2.

    Resolution chain (dept2 first, then dept1):
      1. Exact alias match
      2. Exact match in DB departments
      3. Substring/contains match against DB departments
      4. Fuzzy match (difflib) against DB departments
      5. Fallback: original value or "未知"

    Returns (resolved_department, match_method) where match_method
    describes how the match was made (for logging).
    """
    # ── Try dept2 first (more specific) ──
    if dept2:
        # 1. Alias
        canonical = aliases.get(dept2)
        if canonical and canonical in db_depts:
            return canonical, f"别名匹配(dept2): {dept2} → {canonical}"

        # 2. Exact
        if dept2 in db_depts:
            return dept2, f"精确匹配(dept2): {dept2}"

        # 3. Substring/fuzzy
        match = _fuzzy_match_dept(dept2, db_depts)
        if match:
            return match, f"模糊匹配(dept2): {dept2} → {match}"

    # ── Try dept1 ──
    if dept1:
        # 1. Alias
        canonical = aliases.get(dept1)
        if canonical and canonical in db_depts:
            return canonical, f"别名匹配(dept1): {dept1} → {canonical}"

        # 2. Exact
        if dept1 in db_depts:
            return dept1, f"精确匹配(dept1): {dept1}"

        # 3. Substring/fuzzy
        match = _fuzzy_match_dept(dept1, db_depts)
        if match:
            return match, f"模糊匹配(dept1): {dept1} → {match}"

    # ── Fallback ──
    fallback = dept2 or dept1 or "未知"
    return fallback, f"未匹配: 使用原始值 [{fallback}]"


def import_persons_from_excel(dry_run: bool = False) -> dict:
    """
    Read the Excel file and import all persons into the database.
    Department resolution uses fuzzy matching against the departments table.
    """
    filepath = Path(PEOPLE_DIR) / EXCEL_FILE
    if not filepath.exists():
        return {"error": f"Excel file not found: {filepath}"}

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # Map column indices from headers
    headers = [cell.value for cell in ws[1]]
    col = {}
    for j, h in enumerate(headers):
        if h:
            col[h] = j

    # Load department names from DB for matching
    db_depts = _load_db_departments()
    aliases = DEPT_NAME_ALIASES
    print(f"[Importer] Loaded {len(db_depts)} departments from database.")

    # Track match statistics
    match_stats: dict[str, int] = {}

    # Parse all rows (skip header)
    entries = []
    unmatched: list[dict] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        row_data = list(row)
        name = row_data[col.get("人员姓名", 0)]
        level1 = row_data[col.get("所属部门一级分类", 1)]
        level2 = row_data[col.get("所属二级部门", 2)]
        person_type = row_data[col.get("所属人员类型", 3)]
        pos_level_name = row_data[col.get("职位等级", 4)]
        pos_level_num = row_data[col.get("职位等级(数字化)", 5)]
        phone = row_data[col.get("手机号", 6)]
        email = row_data[col.get("邮箱", 7)]

        if not name:
            continue

        # Clean up values
        dept1 = str(level1).strip() if level1 else ""
        dept2 = str(level2).strip() if level2 else ""
        name = str(name).strip()

        # Resolve department with fuzzy matching
        department, match_method = resolve_department_fuzzy(dept1, dept2, db_depts, aliases)

        # Track statistics
        method_key = match_method.split(":")[0].split("(")[0] if ":" in match_method else match_method
        match_stats[method_key] = match_stats.get(method_key, 0) + 1

        if match_method.startswith("未匹配"):
            unmatched.append({
                "name": name,
                "dept1": dept1,
                "dept2": dept2,
                "resolved": department,
            })

        # Determine person_type
        if person_type and str(person_type).strip() in PERSON_TYPES:
            pt = str(person_type).strip()
        else:
            pt = person_type or "员工"

        # Determine position_title and position_level
        pos_level_name = str(pos_level_name).strip() if pos_level_name else ""
        try:
            pos_num_val = float(pos_level_num) if pos_level_num else 0
        except (ValueError, TypeError):
            pos_num_val = 0

        if pos_level_name and pos_level_name in POSITION_LEVEL_MAP:
            position_title = pos_level_name
            position_level = POSITION_LEVEL_MAP[pos_level_name]
        elif pos_num_val > 0:
            position_level = pos_num_val
            position_title = pos_level_name or None
        else:
            position_title = None
            position_level = None

        # Use position_title or person_type as the "position" field
        position = position_title or pt or ""

        # Generate synthetic filename (will be updated by photo matching later)
        filename = f"{name}-{department}-{position}.jpg"

        entries.append({
            "name": name,
            "department": department,
            "position": position,
            "filename": filename,
            "person_type": pt,
            "position_title": position_title,
            "position_level": position_level,
            "phone": str(phone).strip() if phone else None,
            "email": str(email).strip() if email else None,
        })

    # Assign sort_order within each (department, person_type) group
    groups: dict[tuple, list] = {}
    for entry in entries:
        key = (entry["department"], entry["person_type"])
        groups.setdefault(key, []).append(entry)
    for group_entries in groups.values():
        for idx, entry in enumerate(group_entries):
            entry["sort_order"] = idx

    # Deduplicate filenames (multiple persons can share same name+dept+position)
    seen_filenames: dict[str, int] = {}
    for entry in entries:
        fn = entry["filename"]
        if fn in seen_filenames:
            seen_filenames[fn] += 1
            base, ext = os.path.splitext(fn)
            entry["filename"] = f"{base}-{seen_filenames[fn]}{ext}"
        else:
            seen_filenames[fn] = 0

    # Print summary
    print(f"[Importer] Parsed {len(entries)} persons in {len(groups)} groups.")
    print(f"[Importer] Match statistics: {match_stats}")
    if unmatched:
        print(f"[Importer] Unmatched departments ({len(unmatched)}):")
        for item in unmatched[:10]:
            print(f"  - {item['name']}: dept1=[{item['dept1']}], dept2=[{item['dept2']}] → [{item['resolved']}]")
        if len(unmatched) > 10:
            print(f"  ... and {len(unmatched) - 10} more")

    if dry_run:
        return {
            "total": len(entries),
            "groups": len(groups),
            "match_stats": match_stats,
            "unmatched": unmatched,
        }

    # Clear existing persons and insert new ones
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM persons")
            cur.execute("ALTER TABLE persons AUTO_INCREMENT = 1")
        conn.commit()
        print("[Importer] Cleared all existing persons.")

        inserted = 0
        with conn.cursor() as cur:
            for entry in entries:
                cur.execute(
                    """INSERT INTO persons
                       (name, department, position, filename,
                        person_type, sort_order,
                        position_title, position_level,
                        phone, email)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                    (entry["name"], entry["department"], entry["position"],
                     entry["filename"], entry["person_type"], entry["sort_order"],
                     entry["position_title"], entry["position_level"],
                     entry["phone"], entry["email"]),
                )
                inserted += 1
        conn.commit()

        result = {
            "inserted": inserted,
            "total": len(entries),
            "match_stats": match_stats,
            "unmatched_count": len(unmatched),
            "unmatched": unmatched[:20] if len(unmatched) <= 20 else unmatched[:20],
        }
        print(f"[Importer] Inserted {inserted} persons.")
        return result
    except Exception as e:
        conn.rollback()
        raise
    finally:
        conn.close()


def match_photos_to_persons() -> dict:
    """
    Scan the photos directory, parse filenames, and match photos to persons
    in the database by name. Updates the filename field for matched persons.

    Photo filename format: 姓名-部门-职位.jpg
    """
    if not os.path.isdir(PHOTOS_DIR):
        print(f"[PhotoMatch] Photos directory not found: {PHOTOS_DIR}")
        return {"error": "Photos directory not found", "matched": 0, "total_photos": 0}

    # Scan photos
    photo_files = []
    for fname in os.listdir(PHOTOS_DIR):
        ext = os.path.splitext(fname)[1].lower()
        if ext in IMAGE_EXTENSIONS:
            photo_files.append(fname)

    print(f"[PhotoMatch] Found {len(photo_files)} photos.")

    # Parse each photo filename
    parsed_photos = []
    for fname in photo_files:
        name_without_ext = os.path.splitext(fname)[0]
        parts = name_without_ext.split("-")
        if len(parts) >= 2:
            photo_name = parts[0]
            # Department is everything between first and last part
            photo_dept = "-".join(parts[1:-1]) if len(parts) > 2 else parts[1]
            parsed_photos.append({
                "filename": fname,
                "name": photo_name,
                "department": photo_dept,
            })
        else:
            parsed_photos.append({
                "filename": fname,
                "name": name_without_ext,
                "department": None,
            })

    # Match photos to persons by name
    conn = get_connection()
    try:
        # Get all persons
        with conn.cursor() as cur:
            cur.execute("SELECT id, name, department FROM persons")
            persons = cur.fetchall()

        # Build lookup: name → list of persons
        name_index: dict[str, list[dict]] = {}
        for p in persons:
            name_index.setdefault(p["name"], []).append(p)

        matched = 0
        unmatched_photos = []
        multi_match_photos = []

        with conn.cursor() as cur:
            for photo in parsed_photos:
                candidates = name_index.get(photo["name"], [])

                if len(candidates) == 0:
                    unmatched_photos.append(photo["filename"])
                    continue

                if len(candidates) == 1:
                    # Unique name match — update filename
                    cur.execute(
                        "UPDATE persons SET filename = %s WHERE id = %s",
                        (photo["filename"], candidates[0]["id"]),
                    )
                    matched += 1
                else:
                    # Multiple persons with same name — try department match
                    dept_match = None
                    for c in candidates:
                        if photo["department"] and c["department"] == photo["department"]:
                            dept_match = c
                            break

                    if dept_match:
                        cur.execute(
                            "UPDATE persons SET filename = %s WHERE id = %s",
                            (photo["filename"], dept_match["id"]),
                        )
                        matched += 1
                    else:
                        # Multiple matches, no department match — use first one
                        cur.execute(
                            "UPDATE persons SET filename = %s WHERE id = %s",
                            (photo["filename"], candidates[0]["id"]),
                        )
                        matched += 1
                        multi_match_photos.append({
                            "filename": photo["filename"],
                            "name": photo["name"],
                            "candidates": [c["name"] + "@" + c["department"] for c in candidates],
                        })

        conn.commit()

        result = {
            "matched": matched,
            "total_photos": len(photo_files),
            "unmatched_photos": unmatched_photos,
            "multi_match_photos": multi_match_photos,
        }
        print(f"[PhotoMatch] Matched {matched}/{len(photo_files)} photos to persons.")
        if unmatched_photos:
            print(f"[PhotoMatch] Unmatched photos ({len(unmatched_photos)}):")
            for f in unmatched_photos[:10]:
                print(f"  - {f}")
            if len(unmatched_photos) > 10:
                print(f"  ... and {len(unmatched_photos) - 10} more")
        if multi_match_photos:
            print(f"[PhotoMatch] Multi-name matches ({len(multi_match_photos)}):")
            for m in multi_match_photos[:5]:
                print(f"  - {m['filename']}: candidates={m['candidates']}")
        return result
    except Exception as e:
        conn.rollback()
        raise
    finally:
        conn.close()


def import_company_leaders_from_markdown():
    """Parse the Markdown hierarchy file and import company leadership entries."""
    filepath = Path(PEOPLE_DIR) / MARKDOWN_FILE
    content = filepath.read_text(encoding='utf-8')

    entries = []
    current_dept = None

    for line in content.strip().split('\n'):
        line = line.strip()
        if not line:
            continue

        # Category header: "1.公司领导" etc.
        if re.match(r'^\d+\.', line):
            continue

        # Department line with people: "* 公司领导：龚勃（2.0）..." or "* 办公室：石宗宏3.0..."
        if line.startswith('*'):
            line = line[1:].strip()

            # Check for "：" or ":" separator
            sep = '：' if '：' in line else (':' if ':' in line else None)
            if sep:
                dept_name = line.split(sep)[0].strip()
                people_text = line.split(sep)[1]
                people = _parse_people_from_text(people_text)
            else:
                continue

            # Determine position_title for each person based on level
            for person in people:
                level = person["level"]

                if dept_name == "公司领导":
                    if level == 2.0:
                        position_title = "二级正"
                    elif level == 2.5:
                        position_title = "二级副"
                    elif level == 2.7:
                        position_title = None
                    else:
                        position_title = None
                elif dept_name == "资深总裁":
                    position_title = "资深总裁"
                elif dept_name == "资深副总裁":
                    position_title = "资深副总裁"
                elif dept_name == "资深经理":
                    position_title = "资深经理"
                else:
                    for title, mapped_level in POSITION_LEVEL_MAP.items():
                        if level and mapped_level == level:
                            position_title = title
                            break
                    else:
                        position_title = None

                position_level = level
                position = position_title or ""
                filename = f"{person['name']}-{dept_name}-{position}.jpg"

                entries.append({
                    "name": person["name"],
                    "department": dept_name,
                    "position": position,
                    "filename": filename,
                    "person_type": "部门领导",
                    "sort_order": None,
                    "position_title": position_title,
                    "position_level": position_level,
                    "phone": None,
                    "email": None,
                })

    # Assign sort_order within each (department, person_type) group
    groups = {}
    for entry in entries:
        key = (entry["department"], entry["person_type"])
        groups.setdefault(key, []).append(entry)
    for group_entries in groups.values():
        for idx, entry in enumerate(group_entries):
            entry["sort_order"] = idx

    if not entries:
        return {"inserted": 0, "updated": 0, "total": 0}

    conn = get_connection()
    try:
        inserted = 0
        updated = 0
        with conn.cursor() as cur:
            for entry in entries:
                cur.execute(
                    "SELECT id, filename FROM persons WHERE name = %s AND department = %s",
                    (entry["name"], entry["department"]),
                )
                existing = cur.fetchone()

                if existing:
                    existing_filename = existing["filename"]
                    if existing_filename and existing_filename.startswith(f"{entry['name']}-{entry['department']}-"):
                        new_filename = entry["filename"]
                    else:
                        new_filename = existing_filename

                    cur.execute(
                        """UPDATE persons SET
                           person_type=%s, sort_order=%s,
                           position_title=%s, position_level=%s,
                           position=%s
                           WHERE id=%s""",
                        (entry["person_type"], entry["sort_order"],
                         entry["position_title"], entry["position_level"],
                         entry["position"], existing["id"]),
                    )
                    updated += 1
                else:
                    cur.execute(
                        """INSERT INTO persons
                           (name, department, position, filename,
                            person_type, sort_order,
                            position_title, position_level)
                           VALUES (%s, %s, %s, %s, %s, %s, %s, %s)""",
                        (entry["name"], entry["department"], entry["position"],
                         entry["filename"], entry["person_type"], entry["sort_order"],
                         entry["position_title"], entry["position_level"]),
                    )
                    inserted += 1
        conn.commit()
        return {"inserted": inserted, "updated": updated, "total": len(entries)}
    finally:
        conn.close()


def _parse_people_from_text(text: str) -> list[dict]:
    """Parse people from text like '龚勃（2.0）、马明(2.5)' or '石宗宏3.0、陈嵩3.5'."""
    people = []
    parts = re.split(r'[、,，]', text)
    for part in parts:
        part = part.strip()
        if not part:
            continue

        # Try format: "姓名（等级）" or "姓名(等级)"
        m = re.match(r'(.+?)[（(](\d+\.?\d*)[）)]', part)
        if m:
            people.append({"name": m.group(1).strip(), "level": float(m.group(2))})
            continue

        # Try format: "姓名等级" like "石宗宏3.0"
        m = re.match(r'(.+?)(\d+\.?\d*)$', part)
        if m:
            people.append({"name": m.group(1).strip(), "level": float(m.group(2))})
            continue

        # Just a name without level
        people.append({"name": part, "level": None})

    return people
